"""
Generates sample candidate resumes in PDF and DOCX formats for testing.
"""

from pathlib import Path
import fitz  # PyMuPDF
import docx

OUTPUT_DIR = Path(__file__).resolve().parent

RESUMES = {
    "alex_morgan_ai_senior": {
        "name": "Alex Morgan",
        "email": "alex.morgan@email.com",
        "phone": "+1 (555) 234-5678",
        "title": "Senior AI & Backend Engineer",
        "summary": "Senior Software Engineer with 6 years of experience building enterprise Python microservices, LLM agent workflows using LangChain and LangGraph, and high-performance RAG pipelines with ChromaDB.",
        "skills": [
            "Python (Expert)", "FastAPI", "LangChain", "LangGraph", "ChromaDB",
            "FAISS", "Model Context Protocol (MCP)", "Docker", "AWS", "CI/CD", "PostgreSQL", "REST APIs"
        ],
        "experience": [
            {
                "role": "Senior AI Engineer - CloudTech Solutions",
                "period": "2022 - Present (2.5 years)",
                "details": "Architected multi-agent customer support system using LangGraph and FastAPI. Implemented hybrid RAG with Chroma vector database reducing retrieval latency by 40%. Spearheaded Model Context Protocol (MCP) integrations."
            },
            {
                "role": "Backend Engineer - DataStream Inc.",
                "period": "2019 - 2022 (3.5 years)",
                "details": "Developed Python backend services handling 10M+ daily events. Containerized microservices using Docker and deployed on AWS ECS."
            }
        ],
        "education": "B.S. in Computer Science - University of Washington (2019)",
        "certifications": ["AWS Certified Solutions Architect", "LangChain Developer Certification"]
    },
    "priya_sharma_backend_mid": {
        "name": "Priya Sharma",
        "email": "priya.sharma@email.com",
        "phone": "+1 (555) 876-5432",
        "title": "Backend Python Developer",
        "summary": "Backend developer with 3.5 years of experience specializing in Django, PostgreSQL, and REST API development. Passionate about transitioning into applied AI engineering.",
        "skills": [
            "Python", "Django", "FastAPI (Intermediate)", "PostgreSQL", "Redis",
            "Git", "Docker (Basics)", "Unit Testing", "Pandas", "Scikit-Learn"
        ],
        "experience": [
            {
                "role": "Backend Developer - FinTech Labs",
                "period": "2021 - Present (3.5 years)",
                "details": "Designed relational database schemas and REST APIs using Django REST Framework. Optimized SQL queries improving endpoint throughput by 25%. Built internal data scraping pipelines."
            }
        ],
        "education": "B.Tech in Information Technology - Anna University (2021)",
        "certifications": ["Python Certified Associate Programmer (PCAP)"]
    },
    "marcus_vance_frontend_junior": {
        "name": "Marcus Vance",
        "email": "marcus.vance@email.com",
        "phone": "+1 (555) 345-9876",
        "title": "Junior Frontend Web Developer",
        "summary": "Frontend developer with 1.5 years of experience building modern user interfaces using React, JavaScript, HTML5, and CSS. Basic familiarity with Python and backend scripting.",
        "skills": [
            "JavaScript (ES6+)", "React.js", "HTML5", "CSS3", "Tailwind CSS",
            "Git", "REST APIs", "Python (Basics)", "Figma"
        ],
        "experience": [
            {
                "role": "Junior Frontend Developer - Creative Agency",
                "period": "2023 - Present (1.5 years)",
                "details": "Built responsive landing pages and client dashboards using React and Tailwind CSS. Collaborated with UI/UX designers to implement pixel-perfect user flows."
            }
        ],
        "education": "B.A. in Digital Arts & Design - Portland State University (2023)",
        "certifications": ["Meta Front-End Developer Professional Certificate"]
    }
}


def build_pdf(candidate_data: dict, filename: Path):
    doc = fitz.open()
    page = doc.new_page(width=595, height=842)  # A4
    
    text = f"""{candidate_data['name'].upper()}
{candidate_data['title']} | {candidate_data['email']} | {candidate_data['phone']}
--------------------------------------------------------------------------------

PROFESSIONAL SUMMARY:
{candidate_data['summary']}

CORE SKILLS:
{', '.join(candidate_data['skills'])}

WORK EXPERIENCE:
"""
    for exp in candidate_data['experience']:
        text += f"\n- {exp['role']} ({exp['period']})\n  {exp['details']}\n"
        
    text += f"""
EDUCATION:
{candidate_data['education']}

CERTIFICATIONS:
{', '.join(candidate_data['certifications'])}
"""
    
    rect = fitz.Rect(50, 50, 545, 800)
    page.insert_textbox(rect, text, fontsize=11, fontname="helv", color=(0.1, 0.1, 0.1))
    doc.save(str(filename))
    doc.close()


def build_docx(candidate_data: dict, filename: Path):
    doc = docx.Document()
    doc.add_heading(candidate_data['name'], level=0)
    doc.add_paragraph(f"{candidate_data['title']} | {candidate_data['email']} | {candidate_data['phone']}")
    
    doc.add_heading("Professional Summary", level=1)
    doc.add_paragraph(candidate_data['summary'])
    
    doc.add_heading("Core Skills", level=1)
    doc.add_paragraph(", ".join(candidate_data['skills']))
    
    doc.add_heading("Work Experience", level=1)
    for exp in candidate_data['experience']:
        p = doc.add_paragraph()
        p.add_run(f"{exp['role']} ({exp['period']})\n").bold = True
        p.add_run(exp['details'])
        
    doc.add_heading("Education", level=1)
    doc.add_paragraph(candidate_data['education'])
    
    doc.add_heading("Certifications", level=1)
    doc.add_paragraph(", ".join(candidate_data['certifications']))
    
    doc.save(str(filename))


def main():
    OUTPUT_DIR.mkdir(exist_ok=True)
    # 1. Alex Morgan (PDF)
    build_pdf(RESUMES["alex_morgan_ai_senior"], OUTPUT_DIR / "alex_morgan_resume.pdf")
    # 2. Priya Sharma (DOCX)
    build_docx(RESUMES["priya_sharma_backend_mid"], OUTPUT_DIR / "priya_sharma_resume.docx")
    # 3. Marcus Vance (PDF)
    build_pdf(RESUMES["marcus_vance_frontend_junior"], OUTPUT_DIR / "marcus_vance_resume.pdf")
    print("Successfully generated sample resumes in PDF and DOCX formats!")

if __name__ == "__main__":
    main()
